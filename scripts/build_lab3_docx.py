"""Build ECEN3314 Lab 3 simulation results Word document."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[1]
ROOT = REPO / "labs" / "lab3_results"
FIGS = ROOT / "figures"
OUT = ROOT / "ECEN3314_Lab3_Simulation_Results.docx"


def set_run_font(run, size=11, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return h


def add_para(doc, text, *, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=10, bold=False)
    run.italic = True
    p.paragraph_format.space_after = Pt(12)
    return p


def add_image(doc, path: Path, width_in: float = 6.2):
    if not path.is_file():
        add_para(doc, f"[Missing figure: {path.name}]", bold=True)
        return
    doc.add_picture(str(path), width=Inches(width_in))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("ECEN 3314 — Lab 3")
    set_run_font(r, size=22, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Diode and Rectifier Circuits — Simulation Results")
    set_run_font(r, size=14, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"OrCAD Capture / PSpice + MATLAB\n"
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
        f"Project: orcadcpmtest"
    )
    set_run_font(r, size=10)

    add_heading(doc, "1. Summary", 1)
    add_para(
        doc,
        "This document compiles Section 6 simulation deliverables [S-1]–[S-8] from "
        "ECEN3314 Lab 3 (Diode and Rectifier Circuits). Circuits were entered in OrCAD Capture; "
        "analyses were run with PSpice (psp_cmd); waveforms and I–V curves were plotted in MATLAB. "
        "PSpice Probe UI screenshots were not available through automation, so MATLAB plots are used "
        "for simulation results.",
    )

    add_heading(doc, "2. Schematics (Capture)", 1)
    add_para(
        doc,
        "All four Lab 3 circuits were placed on SCHEMATIC1 / PAGE1. "
        "Overview screenshot from OrCAD Capture:",
    )
    add_image(doc, FIGS / "S_schematics_CaptureWindow.png", width_in=6.4)
    add_caption(
        doc,
        "Figure A. Capture PAGE1 overview — Fig 9 (diode I–V), Fig 10 (Zener), and Fig 11 (peak rectifier, partial).",
    )

    add_para(
        doc,
        "Schematic notes:\n"
        "• [S-1] Fig 9: V1 (VDC) → D1 (1N4148) → R1 = 1 kΩ → GND\n"
        "• [S-3] Fig 10: V2 (VDC) → R2 = 2 kΩ → DZ (Dbreak / MODEL D1N5231, reverse) → GND\n"
        "• [S-5] Fig 11: V3 (VSIN, 3 V, 50 Hz) → D2 (1N4007) → R3 = 150 Ω || C1 = 47 µF → GND\n"
        "• [S-7] Fig 8: Full-wave bridge (4×1N4007) + RL = 10 kΩ, VSIN 3 V / 50 Hz "
        "(on page below/right of overview crop)",
    )

    add_heading(doc, "3. [S-2] Diode I–V Characteristics (Fig 9)", 1)
    add_para(
        doc,
        "DC sweep of VIN from 0 V to 5 V (0.01 V step). Diode 1N4148 with series 1 kΩ. "
        "Plotted: diode current ID vs diode voltage VD = V(anode) − V(cathode).",
    )
    add_image(doc, FIGS / "S2_diode_IV.png", width_in=5.8)
    add_caption(doc, "Figure B. [S-2] 1N4148 forward I–V from PSpice DC sweep (MATLAB).")
    add_para(
        doc,
        "Observed: knee near ~0.5–0.6 V; at VIN = 5 V, VD ≈ 0.67 V and ID ≈ 4.3 mA "
        "(consistent with ~0.7 V silicon forward drop and IR drop on 1 kΩ).",
    )

    add_heading(doc, "4. [S-4] Zener Diode Reverse I–V (Fig 10)", 1)
    add_para(
        doc,
        "DC sweep of VIN from 0 V to 15 V (0.05 V step). 1N5231 (Vz ≈ 5.1 V) reverse-biased "
        "with series 2 kΩ (VIN+ → R → cathode; anode → GND).",
    )
    add_image(doc, FIGS / "S4_zener_IV.png", width_in=5.8)
    add_caption(doc, "Figure C. [S-4] 1N5231 reverse characteristic — Vcathode vs current (MATLAB).")
    add_para(
        doc,
        "Observed: current rises sharply near Vz = 5.1 V. At VIN = 15 V, Vcathode ≈ 5.04 V "
        "and I ≈ 5.0 mA (regulation around the Zener voltage).",
    )

    add_heading(doc, "5. [S-6] Peak Rectifier (Fig 11)", 1)
    add_para(
        doc,
        "Transient analysis, VSIN amplitude 3 V, 0 V offset. Diode 1N4007, R = 150 Ω, C = 47 µF. "
        "Each plot shows two full input cycles with VIN and VOUT.",
    )

    for freq, ripple, vpk in (
        (50, 2.007, 2.274),
        (150, 1.216, 2.265),
        (500, 0.485, 2.236),
    ):
        add_heading(doc, f"5.{1 if freq == 50 else 2 if freq == 150 else 3}  Frequency = {freq} Hz", 2)
        add_image(doc, FIGS / f"S6_peak_{freq}Hz.png", width_in=5.8)
        add_caption(doc, f"Figure D-{freq}. [S-6] Peak rectifier @ {freq} Hz — VIN (blue), VOUT (red).")
        add_para(
            doc,
            f"VIN(peak) = 3.0 V, VO(peak) ≈ {vpk:.3f} V, "
            f"VRipple(p-p) ≈ {ripple:.3f} V. "
            f"Diode drop ≈ VIN(peak) − VO(peak) ≈ {3.0 - vpk:.2f} V.",
        )

    add_para(
        doc,
        "Trend: ripple decreases as frequency increases (less discharge time per cycle for fixed RC). "
        "Large ripple at 50 Hz is expected with R = 150 Ω and C = 47 µF (τ = RC ≈ 7 ms vs T = 20 ms).",
    )

    add_heading(doc, "6. [S-8] Full-Wave Bridge Rectifier (Fig 8)", 1)
    add_para(
        doc,
        "Transient analysis, VSIN 3 V peak, 50 Hz, two cycles. Bridge of four 1N4007 diodes with "
        "RL = 10 kΩ. VOUT measured across the load (VA − VB).",
    )
    add_image(doc, FIGS / "S8_fullwave_50Hz.png", width_in=5.8)
    add_caption(doc, "Figure E. [S-8] Full-wave bridge @ 50 Hz — VIN and rectified VOUT (MATLAB).")
    add_para(
        doc,
        "Observed: output is full-wave rectified with peak ≈ 1.9–2.0 V "
        "(≈ two diode forward drops below the 3 V input peak). Dead zones near zero-crossings "
        "where |VIN| is below 2×Von.",
    )

    add_heading(doc, "7. Simulation Setup Notes", 1)
    add_para(
        doc,
        "• Diode models: labs/lab3_results/models (1N4148, 1N4007, D1N5231).\n"
        "• Headless runs via psp_cmd.exe; .PRINT used V([node]) syntax.\n"
        "• CSV data: labs/lab3_results/data/\n"
        "• Waveform plots: MATLAB Engine MCP (R2024b).",
    )

    add_heading(doc, "8. File Index", 1)
    add_para(
        doc,
        "Report root: labs/lab3_results/\n"
        "figures/ — Capture overview + MATLAB plots\n"
        "data/ — CSV traces\n"
        "netlists/ — .cir used for psp_cmd\n"
        "models/ — diode .lib files",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
